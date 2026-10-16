"""CSV and XLSX parsing helpers for Acollida fee imports."""

import base64
import csv
import io
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation

from odoo import _
from odoo.exceptions import UserError

MONTHS = (
    ('setembre', 9),
    ('octubre', 10),
    ('novembre', 11),
    ('desembre', 12),
    ('gener', 1),
    ('febrer', 2),
    ('marc', 3),
    ('abril', 4),
    ('maig', 5),
    ('juny', 6),
)

_ALIASES = {
    'student': 'student',
    'alumne': 'student',
    'name': 'student',
    'nen': 'student',
    'parent': 'parent',
    'pare': 'parent',
    'course': 'course',
    'curs': 'course',
    'tranche': 'tranche',
    'tram': 'tranche',
    'weekly_hours': 'weekly_hours',
    'hores_setmana': 'weekly_hours',
    'hours_per_week': 'weekly_hours',
    'base_amount': 'base_amount',
    'import_mes': 'base_amount',
    'monthly_amount': 'base_amount',
}
_ALIASES.update({name: f'month_{number}' for name, number in MONTHS})
_REQUIRED = {'student', 'course', 'tranche', 'weekly_hours', 'base_amount'}
_DECIMAL_QUANTUM = Decimal('0.01')


def normalize_header(value):
    text = unicodedata.normalize('NFKD', value or '')
    text = ''.join(char for char in text if not unicodedata.combining(char)).casefold()
    return re.sub(r'[^a-z0-9]+', '_', text).strip('_')


def normalize_text(value):
    text = unicodedata.normalize('NFKD', value or '')
    text = ''.join(char for char in text if not unicodedata.combining(char))
    return ' '.join(text.casefold().split())


def normalize_course(value):
    """Match common Catalan ordinal spellings used by school spreadsheets."""
    words = normalize_text(value).split()
    if not words:
        return ''
    aliases = {
        '1er': '1r',
        '2on': '2n',
        '3er': '3r',
        '4art': '4t',
        '5e': '5e',
        '6e': '6e',
    }
    words[0] = aliases.get(words[0], words[0])
    words = [word for word in words if word not in {'primaria', 'primari'}]
    return ' '.join(words)


def parse_amount(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise TypeError(_('A fee must be a number.'))
    text = str(value).strip().replace('\u00a0', '').replace(' ', '').replace('€', '')
    if ',' in text:
        text = text.replace('.', '').replace(',', '.')
    try:
        amount = Decimal(text).quantize(_DECIMAL_QUANTUM)
    except (InvalidOperation, ValueError) as error:
        raise ValueError(_('The fee "%s" is not a valid amount.') % value) from error
    if not amount.is_finite():
        raise ValueError(_('The fee "%s" is not a valid amount.') % value)
    if amount < 0:
        raise ValueError(_('Fees cannot be negative.'))
    return amount


def parse_hours(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, time):
        return Decimal(value.hour) + Decimal(value.minute) / Decimal(60)
    if isinstance(value, timedelta):
        hours = Decimal(str(value.total_seconds())) / Decimal(3600)
        if hours < 0:
            raise ValueError(_('Weekly hours cannot be negative.'))
        return hours
    text = str(value).strip().replace(',', '.')
    if ':' in text:
        pieces = text.split(':')
        if len(pieces) < 2 or len(pieces) > 3:
            raise ValueError(_('Weekly hours "%s" are not valid.') % value)
        try:
            hours, minutes = int(pieces[0]), int(pieces[1])
            seconds = Decimal(pieces[2]) if len(pieces) == 3 else Decimal(0)
        except (InvalidOperation, ValueError) as error:
            raise ValueError(_('Weekly hours "%s" are not valid.') % value) from error
        if minutes < 0 or minutes >= 60 or not seconds.is_finite() or seconds < 0 or seconds >= 60:
            raise ValueError(_('Weekly hours "%s" are not valid.') % value)
        result = Decimal(hours) + Decimal(minutes) / Decimal(60) + seconds / Decimal(3600)
        if result < 0:
            raise ValueError(_('Weekly hours cannot be negative.'))
        return result
    if '.' in text:
        hours, minute_text = text.split('.', 1)
        if len(minute_text) == 1:
            minute_text += '0'
        if (
            len(minute_text) > 2
            or not minute_text.isdigit()
            or int(minute_text) >= 60
            or not hours.isdigit()
        ):
            raise ValueError(_('Weekly hours "%s" are not valid.') % value)
        return Decimal(int(hours)) + Decimal(int(minute_text)) / Decimal(60)
    try:
        hours = Decimal(text)
    except InvalidOperation as error:
        raise ValueError(_('Weekly hours "%s" are not valid.') % value) from error
    if not hours.is_finite() or hours < 0:
        raise ValueError(_('Weekly hours "%s" are not valid.') % value)
    return hours


def parse_rows(raw, filename):
    filename = (filename or '').strip().lower()
    content = base64.b64decode(raw)
    if filename.endswith('.xlsx'):
        rows = _read_xlsx(content)
    elif filename.endswith(('.csv', '.txt')) or '.' not in filename:
        rows = _read_csv(content)
    else:
        raise UserError(_('Upload a CSV or .xlsx file, not "%s".') % (filename or ''))
    return _records(rows)


def _read_xlsx(content):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise UserError(
            _('Excel support is unavailable in this Odoo installation; upload a CSV instead.')
        ) from error
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        return [
            [_cell_text(value) for value in row]
            for row in workbook.active.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _read_csv(content):
    try:
        text = content.decode('utf-8-sig')
    except UnicodeDecodeError as error:
        raise UserError(_('The CSV file must use UTF-8 encoding.')) from error
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return list(csv.reader(io.StringIO(text), dialect))


def _records(rows):
    header = None
    canonical = []
    records = []
    for row_number, row in enumerate(rows, start=1):
        if not any(str(cell or '').strip() for cell in row):
            continue
        if header is None:
            header = row
            canonical = [_ALIASES.get(normalize_header(cell)) for cell in row]
            recognized = [column for column in canonical if column]
            repeated = sorted({column for column in recognized if recognized.count(column) > 1})
            if repeated:
                raise UserError(_('The file has repeated columns: %s.') % ', '.join(repeated))
            missing = sorted(_REQUIRED - set(recognized))
            missing_months = [
                f'month_{number}' for _name, number in MONTHS if f'month_{number}' not in recognized
            ]
            if missing or missing_months:
                missing_columns = missing + missing_months
                raise UserError(
                    _('The file is missing required columns: %s.') % ', '.join(missing_columns)
                )
            continue
        record = dict.fromkeys(_REQUIRED | {'parent'}, '')
        record.update({f'month_{number}': '' for _name, number in MONTHS})
        for column, value in zip(canonical, row, strict=False):
            if column:
                record[column] = value
        record['row_number'] = row_number
        records.append(record)
    if not records:
        raise UserError(_('The file has no data rows.'))
    return records


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()
