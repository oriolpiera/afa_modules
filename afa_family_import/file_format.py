"""Shared file-format helpers for the AFA family load assistant."""

import base64
import csv
import io
import re
from datetime import date, datetime

from odoo import _
from odoo.exceptions import UserError

COLUMNS = (
    'family_name',
    'role',
    'is_billing',
    'name',
    'email',
    'phone',
    'course',
    'iban',
    'mandate_date',
)

_TEMPLATE_HEADER = ','.join(COLUMNS)
_TEMPLATE_PAYER = (
    'Roure family,guardian,true,Anna Roure,anna@example.cat,+34600000001,,'
    'ES9121000418450200051332,2026-09-01'
)
_TEMPLATE_GUARDIAN = 'Roure family,guardian,false,Jordi Roure,jordi@example.cat,+34600000002,,,'
_TEMPLATE_STUDENT = 'Roure family,student,false,Pol Roure,,,4t Primària,,'

_TEMPLATE_ROWS = (_TEMPLATE_HEADER, _TEMPLATE_PAYER, _TEMPLATE_GUARDIAN, _TEMPLATE_STUDENT)

TEMPLATE_CSV = '\r\n'.join(_TEMPLATE_ROWS) + '\r\n'

_TRUE_VALUES = frozenset({'true', '1', 'yes', 'y', 'si', 'sí', 's', 'x'})
_FALSE_VALUES = frozenset({'false', '0', 'no', 'n', ''})

_EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]{2,}$')
# IBANs are ASCII only; isalnum()/isdigit() accept Unicode digits that int() rejects.
_IBAN_RE = re.compile(r'^[A-Z0-9]+$')

# ISO 13616 IBAN lengths. Unknown countries fall back to the generic bounds.
_IBAN_LENGTHS = {
    'AD': 24,
    'AE': 23,
    'AL': 28,
    'AT': 20,
    'AZ': 28,
    'BA': 20,
    'BE': 16,
    'BG': 22,
    'BH': 22,
    'BR': 29,
    'BY': 28,
    'CH': 21,
    'CR': 22,
    'CY': 28,
    'CZ': 24,
    'DE': 22,
    'DK': 18,
    'DO': 28,
    'EE': 20,
    'EG': 29,
    'ES': 24,
    'FI': 18,
    'FO': 18,
    'FR': 27,
    'GB': 22,
    'GE': 22,
    'GI': 23,
    'GL': 18,
    'GR': 27,
    'GT': 28,
    'HR': 21,
    'HU': 28,
    'IE': 22,
    'IL': 23,
    'IS': 26,
    'IT': 27,
    'JO': 30,
    'KW': 30,
    'KZ': 20,
    'LB': 28,
    'LC': 32,
    'LI': 21,
    'LT': 20,
    'LU': 20,
    'LV': 21,
    'MC': 27,
    'MD': 24,
    'ME': 22,
    'MK': 19,
    'MR': 27,
    'MT': 31,
    'MU': 30,
    'NL': 18,
    'NO': 15,
    'PK': 24,
    'PL': 28,
    'PS': 29,
    'PT': 25,
    'QA': 29,
    'RO': 24,
    'RS': 22,
    'SA': 24,
    'SC': 31,
    'SE': 24,
    'SI': 19,
    'SK': 24,
    'SM': 27,
    'SV': 28,
    'TL': 23,
    'TN': 24,
    'TR': 26,
    'UA': 29,
    'VA': 22,
    'VG': 24,
    'XK': 20,
}


def normalize_header(value):
    return re.sub(r'[^a-z0-9]+', '_', (value or '').strip().lower()).strip('_')


def valid_email(value):
    return bool(_EMAIL_RE.match((value or '').strip()))


def valid_iban(value):
    iban = ''.join((value or '').split()).upper()
    if len(iban) < 15 or len(iban) > 34 or not _IBAN_RE.match(iban):
        return False
    expected = _IBAN_LENGTHS.get(iban[:2])
    if expected and len(iban) != expected:
        return False
    rearranged = iban[4:] + iban[:4]
    digits = ''.join(char if char.isdigit() else str(ord(char) - 55) for char in rearranged)
    return int(digits) % 97 == 1


def parse_bool(value):
    text = (value or '').strip().lower()
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    return None


def parse_rows(raw, filename):
    """Decode a base64 upload into canonical, header-normalised dictionaries."""
    name = (filename or '').strip().lower()
    if name.endswith('.xlsx'):
        return _read_xlsx(raw)
    if name.endswith(('.csv', '.txt')) or '.' not in name:
        return _read_csv(raw)
    raise UserError(_('Upload a CSV or .xlsx file, not "%s".') % (filename or ''))


def _read_csv(raw):
    text = base64.b64decode(raw).decode('utf-8-sig', errors='replace')
    return _read_dicts(csv.reader(io.StringIO(text)))


def _read_xlsx(raw):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise UserError(
            _('Excel support is unavailable in this Odoo installation; upload a CSV instead.')
        ) from error
    workbook = load_workbook(io.BytesIO(base64.b64decode(raw)), read_only=True, data_only=True)
    try:
        rows = [
            [_cell_text(cell) for cell in row]
            for row in workbook.active.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()
    return _read_dicts(rows)


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _read_dicts(rows):
    records = []
    header = None
    for row in rows:
        if header is None:
            header = [normalize_header(cell) for cell in row]
            recognized = set(COLUMNS)
            repeated = sorted(
                {name for name in header if name in recognized and header.count(name) > 1}
            )
            if repeated:
                raise UserError(_('The file has repeated columns: %s.') % ', '.join(repeated))
            missing = [name for name in ('family_name', 'role', 'name') if name not in header]
            if missing:
                raise UserError(
                    _('The file is missing the required columns: %s.') % ', '.join(missing)
                )
            continue
        if not any((cell or '').strip() for cell in row):
            continue
        record = dict.fromkeys(COLUMNS, '')
        for column, cell in zip(header, row, strict=False):
            if column in record:
                record[column] = (cell or '').strip()
        records.append(record)
    if not records:
        raise UserError(_('The file has no data rows.'))
    return records
